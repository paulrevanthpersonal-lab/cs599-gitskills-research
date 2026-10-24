#!/usr/bin/env python3
"""Build the editable checkpoint from the paragraph source. Requires python-docx."""
from pathlib import Path
import argparse,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def hyperlink(p,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
 t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)

def rich(p,text):
 for part in re.split(r'(https?://\S+|\*[^*]+\*)',text):
  if part.startswith(('http://','https://')):
   url=part.rstrip('.,;');hyperlink(p,url)
   if len(url)!=len(part):p.add_run(part[len(url):])
  elif part.startswith('*') and part.endswith('*'):p.add_run(part[1:-1]).italic=True
  elif part:p.add_run(part)

def main():
 root=Path(__file__).resolve().parent.parent
 ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=root/'docs/checkpoint-1-apa.docx');args=ap.parse_args()
 blocks=re.split(r'\n\s*\n',(root/'paper/checkpoint-1.md').read_text().strip())
 title=blocks[0].removeprefix('# ').strip();d=Document();sec=d.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
 sec.header_distance=Inches(.5)
 for name in ['Normal','Title','Heading 1','Heading 2','Header']:
  s=d.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.font.color.rgb=RGBColor(0,0,0)
  f=s.paragraph_format;f.line_spacing=2;f.space_before=Pt(0);f.space_after=Pt(0);f.widow_control=True
  f.first_line_indent=Inches(.5) if name=='Normal' else Inches(0)
  rf=s.element.get_or_add_rPr().rFonts
  for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+k),None)
  for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Times New Roman')
  if name in ['Title','Heading 1','Heading 2']:
   s.font.bold=True;f.keep_with_next=True;f.alignment=WD_ALIGN_PARAGRAPH.LEFT if name=='Heading 2' else WD_ALIGN_PARAGRAPH.CENTER
  if s.element.pPr is not None:
   for el in list(s.element.pPr):
    if el.tag==qn('w:pBdr'):s.element.pPr.remove(el)
 hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');hp._p.append(fld)
 for _ in range(3):d.add_paragraph('')
 d.add_paragraph(title,'Title');d.add_paragraph('')
 for t in ['Paul Revanth Madasu','CS 599 OSS Research Track','Professor Igor Steinmacher']:
  p=d.add_paragraph(t);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
 d.add_page_break();d.add_paragraph(title,'Title');refs=False
 for block in blocks[1:]:
  if block.startswith('## '):
   heading=block[3:].strip()
   if heading=='References':d.add_page_break();refs=True
   d.add_paragraph(heading,'Heading 1')
  elif block.startswith('### '):d.add_paragraph(block[4:].strip(),'Heading 2')
  else:
   p=d.add_paragraph()
   if refs:p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
   rich(p,' '.join(block.splitlines()))
 d.core_properties.title=title;d.core_properties.author='Paul Revanth Madasu';d.core_properties.subject='CS 599 Checkpoint 1 dataset understanding and study design'
 args.output.parent.mkdir(parents=True,exist_ok=True);d.save(args.output)
 assert not d.tables
 print(args.output)
if __name__=='__main__':main()
